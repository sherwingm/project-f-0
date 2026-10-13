"""A backtest's trades CSV -> a per-trade Excel sheet: the option's path while held and after the exit (upside-only
peak, downside-only dip, with dates), holding time, what holding to the last allowed day would have made, and a
summary per arm and variant with the date-clustered t.

    python -m scanner.trade_sheet data/results/S1_v7_trades.csv --last-exit 1 \
        --rule primary="MAIN · SL 50% + 5 days no new high" --rule no_stop="5 days, no stop" \
        --title "S1 v7 backtest" --verdict "..." --out data/results/S1_v7_trade_sheet.xlsx

Option prices are the daily closes in the cached F&O bhavcopies (data/cache/fo_bhavcopy_*.csv); there is no
intraday option history. Numbers only; the hindsight columns are not a test of anything.
"""
from __future__ import annotations

import argparse
import bisect
from collections import defaultdict
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import backtest_s1_v4 as v4
from .build import ROOT
from .option_engine import TICK
from .stats import clustered_t

CACHE = ROOT / "data" / "cache"
REASONS = {"0_stop": "Stop loss 50%", "1_decay": "No new high", "2_t_minus_1": "T-1", "2_t_minus_2": "T-2",
           "contract_changed": "Excluded (corp. action)", "price_adjusted": "Excluded (corp. action)"}


# ---------------------------------------------------------------- the option's path (pure)
def pct(x: float | None, ref: float | None) -> float | None:
    return None if x is None or not ref else round((x / ref - 1) * 100, 2)


def span(points: list[tuple[str, float]], entry: float) -> dict:
    """Upside-only peak and downside-only dip of (date, close) points, % of the entry price. A peak that never got
    above entry is 0 with no date; a dip that never got below entry is 0 with no date."""
    if not points:
        return {"peak": None, "peak_date": None, "dip": None, "dip_date": None, "high": None}
    dh, hi = max(points, key=lambda p: p[1])
    dl, lo = min(points, key=lambda p: p[1])
    up, down = pct(hi, entry), pct(lo, entry)
    return {"peak": max(up, 0.0), "peak_date": dh if up > 0 else None,
            "dip": min(down, 0.0), "dip_date": dl if down < 0 else None, "high": hi}


def sessions_between(sessions: list[str], a: str | None, b: str | None) -> int | None:
    """Sessions after a up to and including b."""
    if not isinstance(a, str) or not isinstance(b, str):
        return None
    return bisect.bisect_right(sessions, b) - bisect.bisect_right(sessions, a)


# ---------------------------------------------------------------- data
def bhav_sessions(cache: Path = CACHE) -> list[str]:
    return sorted(f"{p.stem[-8:-4]}-{p.stem[-4:-2]}-{p.stem[-2:]}" for p in cache.glob("fo_bhavcopy_*.csv"))


def last_day(sessions: list[str], expiry: str, last_exit: int) -> str:
    """The session `last_exit` sessions before expiry (T-1, T-2)."""
    return sessions[bisect.bisect_right(sessions, expiry) - 1 - last_exit]


def option_closes(t: pd.DataFrame, sessions: list[str], cache: Path = CACHE) -> dict:
    """{(date, symbol, expiry, strike): call close} for every trade's sessions entry..last day, one read per file."""
    need = defaultdict(set)
    for r in t.itertuples():
        for d in sessions[bisect.bisect_left(sessions, r.entry_date):bisect.bisect_right(sessions, r.last_day)]:
            need[d].add((r.symbol, r.expiry, float(r.strike)))
    out = {}
    for d in sorted(need):
        b = pd.read_csv(cache / f"fo_bhavcopy_{d.replace('-', '')}.csv",
                        usecols=["TckrSymb", "FinInstrmTp", "XpryDt", "StrkPric", "OptnTp", "ClsPric"])
        b = b[(b["FinInstrmTp"] == "STO") & (b["OptnTp"] == "CE") & b["TckrSymb"].isin({k[0] for k in need[d]})]
        for s, e, k, c in zip(b["TckrSymb"], b["XpryDt"].astype(str).str[:10], b["StrkPric"].astype(float),
                              b["ClsPric"]):
            if (s, e, k) in need[d]:
                out[(d, s, e, k)] = float(c)
    return out


def rows(trades: pd.DataFrame, last_exit: int, cache: Path = CACHE) -> pd.DataFrame:
    t = trades.copy()
    for c in ("signal_date", "entry_date", "exit_date", "expiry"):
        t[c] = t[c].astype(str).str[:10]
    sessions = bhav_sessions(cache)
    t["last_day"] = [last_day(sessions, e, last_exit) for e in t["expiry"]]
    closes = option_closes(t, sessions, cache)
    out = []
    for r in t.itertuples():
        key = (r.symbol, r.expiry, float(r.strike))
        pts = lambda a, b, incl: [(d, closes[(d, *key)]) for d in sessions[   # noqa: E731
            bisect.bisect_left(sessions, a):bisect.bisect_right(sessions, b)] if (incl or d > a) and (d, *key) in closes]
        held = span(pts(r.entry_date, r.exit_date, True), r.entry_price)
        after = span(pts(r.exit_date, r.last_day, False), r.entry_price)
        best = span(pts(r.entry_date, r.last_day, True), r.entry_price)
        last_close = closes.get((r.last_day, *key))
        hold_net = None
        if last_close is not None:
            slip = v4.SLIPPAGE.get(r.category if isinstance(r.category, str) else "unknown", v4.SLIPPAGE["unknown"])
            hold_net = v4.net_of(r.entry_price, max(0.0, round(last_close - TICK, 2)), int(r.qty), slip)[2]
        excluded = r.excluded if isinstance(r.excluded, str) else ""
        n_peak = sessions_between(sessions, r.entry_date, held["peak_date"])
        out.append({
            "arm": r.arm, "variant": r.variant, "symbol": r.symbol, "sector": r.sector or None,
            "results_day": r.signal_date, "beat": r.ar0_pct, "move": r.raw0_pct,
            "near": "Yes" if getattr(r, "near_expiry", False) in (True, "True") else "",
            "kind": {"beat": "Good", "miss": "Bad"}.get(getattr(r, "kind", ""), ""),
            "entry": r.entry_date, "exit": r.exit_date, "held": r.held,
            "days": (pd.Timestamp(r.exit_date) - pd.Timestamp(r.entry_date)).days,
            "reason": REASONS.get(r.exit_reason if isinstance(r.exit_reason, str) else excluded, excluded),
            "expiry": r.expiry, "cost": r.premium_rs, "net_rs": r.net, "net_pct": r.net_pct,
            "h_peak": held["peak"], "h_peak_date": held["peak_date"], "h_peak_day": None if n_peak is None else n_peak + 1,
            "h_dip": held["dip"], "h_dip_date": held["dip_date"],
            "a_peak": after["peak"], "a_peak_date": after["peak_date"],
            "a_peak_n": sessions_between(sessions, r.exit_date, after["peak_date"]),
            "a_dip": after["dip"], "a_dip_date": after["dip_date"],
            "above_exit": None if after["high"] is None or r.exit_price != r.exit_price else after["high"] > r.exit_price,
            "above_entry": None if after["high"] is None else after["high"] > r.entry_price,
            "last_rs": hold_net, "last_pct": None if hold_net is None else round(hold_net / r.premium_rs * 100, 2),
            "best": best["peak"], "best_date": best["peak_date"], "excluded": excluded})
    return pd.DataFrame(out)


# ---------------------------------------------------------------- the workbook
def columns(last_exit: int, near: bool, kinds: bool = False) -> list[tuple]:
    """(group, header, field, kind, width)."""
    after, last = f"OPTION AFTER EXIT (to T-{last_exit})", f"IF HELD TO T-{last_exit}"
    return [
        ("TRADE", "Arm", "arm", "text", 5), ("TRADE", "Symbol", "symbol", "bold", 13),
        ("TRADE", "Sector", "sector", "text", 18), ("TRADE", "Results day", "results_day", "date", 10),
        *([("TRADE", "Result", "kind", "text", 6)] if kinds else []),
        ("TRADE", "Beat vs market", "beat", "pct", 8), ("TRADE", "Results-day move", "move", "pct", 8),
        *([("HOLDING", "Near expiry", "near", "text", 7)] if near else []),
        ("HOLDING", "Entry", "entry", "date", 10), ("HOLDING", "Exit", "exit", "date", 10),
        ("HOLDING", "Sessions held", "held", "int", 8), ("HOLDING", "Calendar days", "days", "int", 8),
        ("HOLDING", "Exit reason", "reason", "text", 15), ("HOLDING", "Expiry", "expiry", "date", 10),
        ("RESULT", "1 lot cost ₹", "cost", "rs", 9), ("RESULT", "Net P&L ₹", "net_rs", "rs_signed", 10),
        ("RESULT", "Net P&L", "net_pct", "pct", 9),
        ("OPTION WHILE HELD", "Max peak", "h_peak", "pct", 9), ("OPTION WHILE HELD", "Peak date", "h_peak_date", "date", 10),
        ("OPTION WHILE HELD", "Peak day (1 = entry)", "h_peak_day", "int", 8),
        ("OPTION WHILE HELD", "Max dip", "h_dip", "pct", 9), ("OPTION WHILE HELD", "Dip date", "h_dip_date", "date", 10),
        (after, "Max peak", "a_peak", "pct", 9), (after, "Peak date", "a_peak_date", "date", 10),
        (after, "Sessions after exit", "a_peak_n", "int", 8), (after, "Max dip", "a_dip", "pct", 9),
        (after, "Dip date", "a_dip_date", "date", 10),
        (last, "Net ₹", "last_rs", "rs_signed", 10), (last, "Net", "last_pct", "pct", 9),
        ("BEST EXIT (hindsight)", "Best", "best", "pct", 9), ("BEST EXIT (hindsight)", "Date", "best_date", "date", 10),
    ]


GROUP_COLOR = {"TRADE": "334155", "HOLDING": "475569", "RESULT": "0F766E", "OPTION WHILE HELD": "1D4ED8",
               "OPTION AFTER": "7C3AED", "IF HELD": "B45309", "BEST EXIT (hindsight)": "6B7280"}
SUB_COLOR = {"TRADE": "E2E8F0", "HOLDING": "E2E8F0", "RESULT": "CCFBF1", "OPTION WHILE HELD": "DBEAFE",
             "OPTION AFTER": "EDE9FE", "IF HELD": "FEF3C7", "BEST EXIT (hindsight)": "F3F4F6"}
FMT = {"pct": '+0.0"%";-0.0"%";0.0"%"', "rs": "#,##0", "rs_signed": "+#,##0;-#,##0;0", "int": "0",
       "date": "dd-mmm-yy", "pct0": '0"%"', "dec": "0.0", "t": "+0.00;-0.00;0.00"}
BASE, BOLD = Font(name="Calibri", size=10), Font(name="Calibri", size=10, bold=True)
STRIPE = PatternFill("solid", fgColor="F8FAFC")
THIN, EDGE = Side(style="thin", color="CBD5E1"), Side(style="medium", color="94A3B8")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _color(table: dict, group: str) -> str:
    return next(v for k, v in table.items() if group.startswith(k))


def _signed(v, bold=False) -> Font:
    return Font(name="Calibri", size=10, bold=bold, color="15803D" if v > 0 else "B91C1C" if v < 0 else "000000")


def trades_sheet(wb, title: str, df: pd.DataFrame, cols: list[tuple]) -> None:
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    start = 1
    for i, c in enumerate(cols, 1):
        if i == len(cols) or cols[i][0] != c[0]:                     # last column of its group
            ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=i)
            cell = ws.cell(row=1, column=start, value=c[0])
            cell.font, cell.alignment = Font(bold=True, color="FFFFFF", size=10), CENTER
            for j in range(start, i + 1):
                ws.cell(row=1, column=j).fill = PatternFill("solid", fgColor=_color(GROUP_COLOR, c[0]))
            start = i + 1
        h = ws.cell(row=2, column=i, value=c[1])
        h.font, h.alignment = Font(bold=True, size=9, color="1E293B"), CENTER
        h.fill = PatternFill("solid", fgColor=_color(SUB_COLOR, c[0]))
        h.border = Border(bottom=EDGE)
        ws.column_dimensions[get_column_letter(i)].width = c[4] + 1
    ws.row_dimensions[1].height, ws.row_dimensions[2].height = 20, 30
    ends = {i for i, c in enumerate(cols, 1) if i == len(cols) or cols[i][0] != c[0]}
    for r, rec in enumerate(df.to_dict("records"), 3):
        for i, (_, _, field, kind, _) in enumerate(cols, 1):
            v = rec[field]
            if v is None or (isinstance(v, float) and v != v):
                v = None
            elif kind == "date":
                v = pd.Timestamp(v).to_pydatetime()
            elif kind == "int":
                v = int(v)
            cell = ws.cell(row=r, column=i, value=v)
            cell.font = BOLD if kind == "bold" else _signed(v) if kind in ("pct", "rs_signed") and v is not None \
                else BASE
            if kind in FMT:
                cell.number_format = FMT[kind]
            cell.alignment = Alignment(horizontal="left" if kind in ("text", "bold") else "right", vertical="center")
            if r % 2 == 0:
                cell.fill = STRIPE
            cell.border = Border(bottom=THIN, right=EDGE if i in ends else None)
    ws.freeze_panes = ws.cell(row=3, column=3 if cols[0][2] != "symbol" else 2)
    ws.auto_filter.ref = f"A2:{get_column_letter(len(cols))}{ws.max_row}"


def measures(g: pd.DataFrame, last_exit: int, near: bool) -> list[tuple]:
    ok = g[g["excluded"] == ""]
    m = lambda c: ok[c].astype(float).mean()                         # noqa: E731
    share = lambda mask: mask.mean() * 100 if len(ok) else None      # noqa: E731
    out = [("Trades", len(ok), "int"),
           *([("  of which near-expiry", (ok["near"] == "Yes").sum(), "int"),
              ("  near-expiry: average net", ok[ok["near"] == "Yes"]["net_pct"].astype(float).mean(), "pct"),
              ("  normal: average net", ok[ok["near"] != "Yes"]["net_pct"].astype(float).mean(), "pct")]
             if near else []),
           ("Winning trades", share(ok["net_rs"] > 0), "pct0"),
           ("Average net per trade", m("net_pct"), "pct"),
           ("t (needs ≥ 2 to count as real)", clustered_t(ok["net_pct"].astype(float), ok["results_day"])
            if len(ok) > 1 else None, "t"),
           ("2021–22 average net", ok[ok["entry"] < "2023"]["net_pct"].astype(float).mean(), "pct"),
           ("2023–24 average net", ok[ok["entry"] >= "2023"]["net_pct"].astype(float).mean(), "pct"),
           ("Total net ₹ (1 lot each)", ok["net_rs"].astype(float).sum(), "rs"),
           ("Average sessions held", m("held"), "dec"), ("Average calendar days held", m("days"), "dec")]
    for label, code in (("Stop loss", "Stop loss 50%"), ("No new high", "No new high"),
                        (f"T-{last_exit}", f"T-{last_exit}")):
        out.append((f"Exit: {label}", share(ok["reason"] == code), "pct0"))
    return out + [
        ("Option while held: avg max peak", m("h_peak"), "pct"), ("Option while held: avg max dip", m("h_dip"), "pct"),
        ("Option after exit: avg max peak", m("a_peak"), "pct"), ("Option after exit: avg max dip", m("a_dip"), "pct"),
        ("Option went above our exit price after we sold", share(ok["above_exit"] == True), "pct0"),   # noqa: E712
        ("Option went above our entry price after we sold", share(ok["above_entry"] == True), "pct0"),  # noqa: E712
        (f"If held to T-{last_exit}: average net", m("last_pct"), "pct"),
        (f"If held to T-{last_exit}: winning trades", share(ok["last_rs"].astype(float) > 0), "pct0"),
        (f"If held to T-{last_exit}: total net ₹", ok["last_rs"].astype(float).sum(), "rs"),
        ("Best exit with hindsight: median", ok["best"].astype(float).median(), "pct")]


ARM_ORDER = ("S", "S_miss", "A", "B")


def arms_of(s: pd.DataFrame) -> list[str]:
    present = list(dict.fromkeys(s["arm"].astype(str)))
    return [a for a in ARM_ORDER if a in present] + [a for a in present if a not in ARM_ORDER]


def summary_sheet(wb, s: pd.DataFrame, rules: dict, title: str, verdict: str, notes: list[str], last_exit: int,
                  near: bool) -> None:
    ws = wb.active
    ws.title = "Summary"
    ws.sheet_view.showGridLines = False
    ws["A1"], ws["A2"] = title, verdict
    ws["A1"].font, ws["A2"].font = Font(bold=True, size=14, color="0F172A"), Font(size=10, color="475569")
    ws.column_dimensions["A"].width = 46
    main = next(iter(rules))
    arms = arms_of(s)
    blocks = [(rule, arm) for rule in rules for arm in arms]
    for j, (rule, arm) in enumerate(blocks, 2):
        ws.column_dimensions[get_column_letter(j)].width = 15
        if arm == arms[0]:
            if len(arms) > 1:
                ws.merge_cells(start_row=4, start_column=j, end_row=4, end_column=j + len(arms) - 1)
            top = ws.cell(row=4, column=j, value=rules[rule])
            top.font, top.alignment = Font(bold=True, color="FFFFFF", size=10), CENTER
            top.fill = PatternFill("solid", fgColor="1D4ED8" if rule == main else "64748B")
        h = ws.cell(row=5, column=j, value=f"Arm {arm}")
        h.font, h.alignment = Font(bold=True, size=10), CENTER
        h.fill = PatternFill("solid", fgColor="DBEAFE" if rule == main else "F1F5F9")
    ws.row_dimensions[4].height = 30
    ws.cell(row=5, column=1, value="Measure").font = Font(bold=True, size=10)
    for j, (rule, arm) in enumerate(blocks, 2):
        for i, (name, v, kind) in enumerate(measures(s[(s["arm"] == arm) & (s["variant"] == rule)], last_exit, near), 6):
            ws.cell(row=i, column=1, value=name).font = BASE
            v = None if v is None or v != v else float(v)
            c = ws.cell(row=i, column=j, value=v)
            c.number_format, c.alignment = FMT[kind], Alignment(horizontal="right")
            c.font = _signed(v, rule == main) if kind in ("pct", "rs", "t") and v is not None else \
                (BOLD if rule == main else BASE)
            for col in (1, j):
                ws.cell(row=i, column=col).border = Border(bottom=THIN)
                if i % 2 == 0:
                    ws.cell(row=i, column=col).fill = STRIPE
    r = ws.max_row + 2
    for k, line in enumerate(notes):
        ws.cell(row=r + k, column=1, value=line).font = Font(bold=k == 0, size=11 if k == 0 else 10, color="334155")
    ws.freeze_panes = "B6"


def default_notes(last_exit: int) -> list[str]:
    return [
        "How to read",
        "• Max peak = how far ABOVE the price we paid the option closed at its highest (0 % = never above entry). "
        "Max dip = how far BELOW the price we paid it closed at its lowest (0 % = never below entry). Daily closes.",
        f"• 'After exit' runs from the session after we sold up to T-{last_exit}, the last day the rules may hold "
        f"(stock options settle by delivery). Blank when the trade already exited on that day.",
        "• 'Peak day' = which day of the trade the peak came (1 = entry day); 'Sessions after exit' = how long after "
        "we sold.",
        "• Net P&L is after brokerage, taxes, exchange charges and slippage (1 % large / 2 % mid / 3 % small cap per "
        "side).",
        f"• 'If held to T-{last_exit}' and 'Best exit' are HINDSIGHT on the same data: ideas for a next version, not "
        "proof.",
        "• Arm A = every results beat. Arm B = only when the stock's sector index is also strong."]


def build(trades: pd.DataFrame, out: Path, rules: dict, last_exit: int, title: str, verdict: str,
          notes: list[str] | None = None) -> pd.DataFrame:
    trades = trades[trades["variant"].isin(rules)]
    s = rows(trades, last_exit)
    near = bool((s["near"] == "Yes").any())
    cols = columns(last_exit, near, kinds=bool((s["kind"] != "").any()))
    wb = Workbook()
    summary_sheet(wb, s, rules, title, verdict, notes or default_notes(last_exit), last_exit, near)
    main = s[s["variant"] == next(iter(rules))].sort_values("entry")
    for arm in arms_of(s):
        trades_sheet(wb, f"Arm {arm} trades", main[main["arm"] == arm], [c for c in cols if c[2] != "arm"])
    order = {k: i for i, k in enumerate(rules)}
    allr = s.assign(rule=s["variant"].map(rules), o=s["variant"].map(order)).sort_values(["o", "arm", "entry"])
    trades_sheet(wb, "All variants", allr, [("TRADE", "Variant", "rule", "text", 22)] + cols)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    s.to_csv(out.with_suffix(".csv"), index=False)
    return s


def main() -> None:
    ap = argparse.ArgumentParser(description="Per-trade Excel sheet of a backtest's trades CSV")
    ap.add_argument("trades", type=Path)
    ap.add_argument("--out", type=Path, required=True)
    ap.add_argument("--last-exit", type=int, required=True, help="sessions before expiry of the last exit (T-n)")
    ap.add_argument("--rule", action="append", required=True, help='variant="label"; the first is the main one')
    ap.add_argument("--title", required=True)
    ap.add_argument("--verdict", default="")
    ap.add_argument("--note", action="append", help="replaces the default notes (first line = heading)")
    a = ap.parse_args()
    rules = dict(r.split("=", 1) for r in a.rule)
    build(pd.read_csv(a.trades), a.out, rules, a.last_exit, a.title, a.verdict, a.note)
    print(a.out)


if __name__ == "__main__":
    main()
