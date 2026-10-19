"""S1 v10 fade · profit-target exit on day 0, 2025-26 (numbers only).

    python -m scanner.fade_target_exit

strategies/S1_v10_backtest_2025_26.md, addendum 2026-10-19 (2). The floor sweep's one-tick trades (entry day 0's open +
1 tick), at the Rs 2 and Rs 1 floors. Exit: a resting sell at entry x (1 + target), rounded up to the tick, filled when
day 0's option high >= target + 1 tick, else day 0's close - 1 tick (E1). Writes data/results/S1_v10_target_exit.md and
S1_v10_target_exit.xlsx.
"""
from __future__ import annotations

import math

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import backtest_s1_v4 as v4
from .backtest_s1_v2 import md
from .backtest_s1_v10 import TICK, net_of
from .build import ROOT
from .fade_floor_sheet import ohlc
from .stats import clustered_t

RES = ROOT / "data" / "results"
TARGETS = (("E1 15:20", None), ("T25", 0.25), ("T50", 0.50), ("T100", 1.00), ("T200", 2.00))
FLOORS = (2.0, 1.0)


def target_price(entry: float, t: float) -> float:
    return round(math.ceil(round(entry * (1 + t) / TICK, 6)) * TICK, 2)


def simulate(t: pd.DataFrame) -> pd.DataFrame:
    """One row per trade and exit: the exit price, whether the target filled, net Rs and net % of premium."""
    rows = []
    for r in t.itertuples(index=False):
        slip = v4.SLIPPAGE.get(r.category, v4.SLIPPAGE["unknown"])
        for name, tg in TARGETS:
            hit = tg is not None and r.h >= target_price(r.entry_price, tg) + TICK - 1e-9
            px = target_price(r.entry_price, tg) if hit else r.e1_exit_price
            _, _, net = net_of(r.side, r.entry_price, px, int(r.lot), slip)
            rows.append({"symbol": r.symbol, "day0": r.day0, "side": r.side, "period": r.period, "fill": r.entry_price,
                         "lot": r.lot, "exit": name, "exit_price": px, "hit": hit, "net": net,
                         "net_pct": round(net / (r.entry_price * r.lot) * 100, 3)})
    return pd.DataFrame(rows)


def stats(g: pd.DataFrame) -> dict:
    r = g["net_pct"]
    return {"n": len(g), "hit_pct": round((g["net"] > 0).mean() * 100, 1), "mean_pct": round(r.mean(), 2),
            "median_pct": round(r.median(), 2), "t": clustered_t(r, g["day0"]) if len(g) > 1 else None,
            "total_rs": round(g["net"].sum()), "target_reached_pct": round(g["hit"].mean() * 100, 1)}


def table(s: pd.DataFrame, by: str | None = None, drop_top: bool = False) -> pd.DataFrame:
    out = []
    for f in FLOORS:
        for name, _ in TARGETS:
            g0 = s[(s["fill"] >= f) & (s["exit"] == name)]
            for key, g in (g0.groupby(by) if by else [("all", g0)]):
                if drop_top and len(g):
                    g = g.drop(index=g["net"].idxmax())
                out.append({"floor_rs": f"{f:g}", "exit": name, **({by: key} if by else {}), **stats(g)})
    return pd.DataFrame(out)


def workbook(t: pd.DataFrame, s: pd.DataFrame, report: dict[str, pd.DataFrame]) -> None:
    F = "Arial"
    head, hfill = Font(name=F, bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F3864")
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    ws["A1"] = "S1 v10 fade · profit-target exit on day 0, 2025-01-01 to 2026-09-25"
    ws["A1"].font = Font(name=F, bold=True, size=14)
    ws["A2"] = ("Values computed in Python (scanner/fade_target_exit.py). Entry: day 0's option open + 1 tick. Target: "
                "a sell at entry x (1 + target), filled if the day's high reached target + 1 tick, else the close - 1 "
                "tick. Net after charges and slippage. Daily data cannot say when the high came, and the live buy is "
                "minutes after the open at the ask, so this is optimistic.")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:J2")
    ws.row_dimensions[2].height = 60
    row = 4
    for title, df in report.items():
        ws.cell(row, 1, title).font = Font(name=F, bold=True, size=12)
        row += 1
        for j, h in enumerate(df.columns, 1):
            c = ws.cell(row, j, h)
            c.font, c.fill = head, hfill
        for rec in df.itertuples(index=False):
            row += 1
            for j, v in enumerate(rec, 1):
                ws.cell(row, j, v).font = Font(name=F)
        row += 2
    for j in range(1, 11):
        ws.column_dimensions[get_column_letter(j)].width = 15
    ws.column_dimensions["A"].width = 10

    tr = wb.create_sheet("Trades")
    wide = s.pivot_table(index=["symbol", "day0", "side"], columns="exit", values=["exit_price", "net"], aggfunc="first")
    wide.columns = [f"{e} {'exit' if v == 'exit_price' else 'net Rs'}" for v, e in wide.columns]
    base = t.set_index(["symbol", "day0", "side"])[["period", "strike", "lot", "o", "h", "l", "c", "entry_price"]]
    base.columns = ["Period", "Strike", "Lot", "Opt open", "Opt high", "Opt low", "Opt close", "Entry"]
    order = [f"{e} {k}" for e, _ in TARGETS for k in ("exit", "net Rs")]
    df = base.join(wide[order]).reset_index().rename(columns={"symbol": "Symbol", "day0": "Day 0", "side": "Side"})
    df.insert(len(df.columns), "High / entry", None)
    tr.append(list(df.columns))
    hi, en = get_column_letter(list(df.columns).index("Opt high") + 1), get_column_letter(list(df.columns).index("Entry") + 1)
    for i, rec in enumerate(df.itertuples(index=False), 2):
        tr.append([None if (isinstance(v, float) and v != v) else v for v in rec][:-1] + [f"={hi}{i}/{en}{i}"])
    for c in tr[1]:
        c.font, c.fill, c.alignment = head, hfill, Alignment(wrap_text=True)
    for j, h in enumerate(df.columns, 1):
        tr.column_dimensions[get_column_letter(j)].width = 11
        fmt = "#,##0;[Red]-#,##0" if "net Rs" in h else "0.00x" if h == "High / entry" else None
        for c in tr.iter_cols(min_col=j, max_col=j, min_row=2):
            for cell in c:
                cell.font = Font(name=F)
                if fmt:
                    cell.number_format = fmt
    tr.freeze_panes = "D2"
    tr.auto_filter.ref = f"A1:{get_column_letter(len(df.columns))}{len(df) + 1}"
    wb.calculation.fullCalcOnLoad = True
    wb.save(RES / "S1_v10_target_exit.xlsx")


def main() -> None:
    t = pd.read_csv(RES / "S1_v10_trades_floor0p05.csv")
    t = t[(t["entry"] == "open") & (t["e1_excluded"].fillna("") == "")].copy().reset_index(drop=True)
    t = t.join(ohlc(t)).sort_values(["day0", "symbol"]).reset_index(drop=True)
    s = simulate(t)
    e1 = s[s["exit"] == "E1 15:20"].set_index(["symbol", "day0", "side"])["net"]
    ref = t.set_index(["symbol", "day0", "side"])["e1_net"]
    diff = (e1 - ref.reindex(e1.index)).abs().max()
    assert diff < 0.011, f"E1 does not reproduce the backtest (max diff Rs {diff})"
    report = {
        "All trades": table(s),
        "Without each row's largest Rs winner": table(s, drop_top=True),
        "Puts and calls": table(s, "side"),
        "By half-year": table(s, "period"),
    }
    text = "\n".join(
        ["# S1 v10 fade · profit-target exit on day 0, 2025-01-01 .. 2026-09-25\n",
         "Numbers only (strategies/S1_v10_backtest_2025_26.md, addendum 2026-10-19 (2)). E1 reproduces the backtest's "
         f"same-day net on all {len(t)} trades (max diff Rs {diff:.2f}). Entry day 0's open + 1 tick; target filled if the "
         "day's high >= target + 1 tick, else the close - 1 tick. Net % of premium after charges and slippage; t "
         "date-clustered. Optimistic: daily data cannot say when the high came, and the live buy is later than the open. "
         "Not out of sample; five exits on one window.\n"]
        + [f"## {k}\n\n{md(v)}" for k, v in report.items()])
    (RES / "S1_v10_target_exit.md").write_text(text, encoding="utf-8")
    workbook(t, s, report)
    print(text)


if __name__ == "__main__":
    main()
