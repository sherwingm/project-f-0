"""Excel writer for scanner.fade_ml: Summary first, then the comparisons, what each model learned, and every trade."""
from __future__ import annotations

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from .fade_indicators_xlsx import F, FB, sheet


def verdict(df: pd.DataFrame) -> str:
    r = df.set_index("method")
    rule = r.iloc[0]
    out = []
    for m in ("Logistic regression", "Gradient boosting", "Small decision tree"):
        a = r.loc[m]
        better = (a["AUC"] or 0) > (rule["AUC"] or 0) and a["mean net % kept"] > rule["mean net % kept"]
        out.append(f"{m}: AUC {a['AUC']:.2f} (rule {rule['AUC']:.2f}, coin flip 0.50); skipping its {int(a['flagged trades'])} riskiest "
                   f"trades keeps {a['mean net % kept']:+.2f} % a trade (rule {rule['mean net % kept']:+.2f} %, all {a['mean net % all']:+.2f} %)"
                   f" → {'beats the rule' if better else 'does NOT beat the rule'}")
    return "\n".join(out)


def write(path, halves, wf, w, txt, imp, per, feats) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    lines = [
        ("S1 v12 fade · can a model learn which trades fail?", FB),
        ("339 backtest trades, v12 exit, real charges. Fail = net < 0. Three models (logistic regression, gradient boosting, one small "
         f"decision tree) learn from {len(feats)} things known before the 09:16 buy, and are scored only on trades they never saw.", F),
        ("AUC = how well the model ranks failures above wins: 0.50 is a coin flip, 1.00 perfect, 0.60 is weak but real. 'Skip flagged' = "
         "drop the model's riskiest trades (as many as the Supertrend-down + gap-down rule drops) and see what the rest earn.", F),
        ("", F),
        ("Test 1 · train on the older half (before 3 Nov 2025), test on the newer half", FB),
        (verdict(halves), F),
        ("", F),
        (f"Test 2 · walk-forward (each quarter predicted by a model trained on every earlier trade; the first quarters had "
         f"too few trades to train on): {wf['test'].iloc[0].replace('walk-forward ', '')}", FB),
        (verdict(wf), F),
        ("", F),
        ("Caution: the rules were found on all 339 trades, the models only on earlier ones, so the rules have a head start. With ~150 "
         "trades to learn from, a few points of AUC either way is noise. Research only: nothing live changes.", F),
        ("Sheets: Test 1 halves · Test 2 walk-forward · Logistic weights · Tree rules · Boosting importance · Trades (each trade's "
         "fail chance from every model, for your own checks).", F),
    ]
    for i, (txt_, font) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=txt_)
        c.font = font
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if "\n" in txt_:
            ws.row_dimensions[i].height = 15 * (txt_.count("\n") + 1) * 1.3
    ws.column_dimensions["A"].width = 160
    cols = ["method", "trades", "AUC", "Brier", "Brier of guessing the fail rate", "flagged trades", "fail % flagged", "fail % kept",
            "fail % all", "mean net % flagged", "mean net % kept", "mean net % all", "total Rs kept (skip flagged)", "total Rs all"]
    sheet(wb, "Test 1 halves", halves[cols], "Trained on the older half, scored on the newer half. Lower Brier = better probabilities "
          "(compare with guessing the fail rate). Rules flag their own trades; models flag the same number of riskiest trades.")
    sheet(wb, "Test 2 walk-forward", wf[cols], "Each quarter predicted by models trained on all earlier trades only.")
    sheet(wb, "Logistic weights", w, "Logistic regression trained on the older half. Inputs are scaled, so weights compare with each other. "
          "Weight 0 = the penalty switched that input off (it did not help).")
    ws2 = wb.create_sheet("Tree rules")
    ws2["A1"] = "The small tree trained on the older half. 'weights: [wins, fails]' = trades in that box. Read it top-down like yes/no questions."
    ws2["A1"].font = Font(name="Arial", size=10, italic=True)
    for i, line in enumerate(txt, 3):
        ws2.cell(row=i, column=1, value=line).font = Font(name="Courier New", size=10)
    ws2.column_dimensions["A"].width = 110
    sheet(wb, "Boosting importance", imp, "Gradient boosting (trained on the older half): how much test-half AUC drops when one input "
          "is shuffled. Near 0 or negative = the model does not really use it.")
    sheet(wb, "Trades", per, "Every trade, the two rules' flags and each model's fail chance (halves test: newer half only; "
          "walk-forward: from the first predicted quarter).")
    wb.save(path)
