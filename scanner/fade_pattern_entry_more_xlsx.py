"""Excel writer for scanner.fade_pattern_entry_more: Summary first, the pass ladder, every rule, the signs, the half-pick
test, and the top rules trade by trade."""
from __future__ import annotations

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, PatternFill

from .fade_indicators_xlsx import GOOD, F, FB, sheet

NEAR = PatternFill("solid", fgColor="FFF2CC")
WITH = " (up for a call, down for a put)"
MEANS = {
    "opt RSI > 50": "Option's RSI 14 on 1-minute candles above 50.",
    "opt RSI > 60": "Option's RSI 14 above 60: a strong push up.",
    "opt RSI crossed above 50": "Option's RSI just moved from 50 or below to above 50.",
    "opt RSI back above 30 (after < 30)": "Option's RSI fell below 30 (oversold) earlier today and is now back above 30.",
    "opt RSI back above 40 (after < 30)": "The same, but waits for the RSI to recover to 40.",
    "opt EMA 9 above EMA 21": "Option's fast EMA (9) above its slow EMA (21): short-term up trend.",
    "opt EMA 9 just crossed above EMA 21": "Option's fast EMA just crossed above the slow one.",
    "opt above VWAP": "Option above its volume-weighted average price since 09:15.",
    "opt just crossed above VWAP": "Option just closed back above its VWAP.",
    "opt MACD above signal": "Option's MACD (12, 26, 9) above its signal line.",
    "opt MACD just crossed above signal": "Option's MACD just crossed above its signal line.",
    "opt back above lower Bollinger (after a close below)": "Option closed below its lower Bollinger band (20, 2) earlier today and is now back above it.",
    "opt above middle Bollinger": "Option's close above its 20-minute average.",
    "opt Stochastic turned up from below 20": "Option's Stochastic %K (14, 3) crossed above %D while below 20 (an oversold turn).",
    "opt volume spike (2× last 20 min) on a green candle": "This minute's volume is more than twice the last 20 minutes' average, and the candle is green.",
    "opt OI lower than 15 candles ago": "Option's open interest is falling: sellers are closing.",
    "opt new high of the day": "Option closed above every earlier high today.",
    "opt above the 09:15 – 09:29 range high": "Option closed above the high of the first 15 minutes (opening-range breakout).",
    "opt 2 green candles in a row": "Two green 1-minute candles in a row.",
    "opt now below the 09:16 buy": "The option is cheaper now than the 09:16 buy price.",
    "VIX above its 09:15 open": "India VIX higher than at 09:15 (fear rising, which lifts option prices).",
    "VIX up vs 15 minutes ago": "India VIX higher than 15 minutes ago.",
}


def means(name: str) -> str:
    if name in MEANS:
        return MEANS[name]
    if name.startswith("opt above EMA"):
        n = name.split()[-1]
        return f"Option's close above its {n}-minute EMA." + (" Before ~11:00 / 12:30 this leans on the day's first prices." if n in ("100", "200") else "")
    if name.startswith("opt Supertrend"):
        n, m = name.split()[2].split("×")
        return f"Option's Supertrend (ATR {n}, × {m}) " + ("just flipped to up." if "turned" in name else "is in up mode.")
    if name.startswith("opt dipped"):
        return f"The option has already traded {name.split()[2]} % below the 09:16 buy price today (the shake-out has happened)."
    if name.startswith("clock"):
        return f"No indicator: just wait until {name.split()[1]}."
    who = "the stock" if name.startswith("stock") else "Nifty 50"
    if "RSI" in name:
        return f"RSI 14 of {who} (1-minute) above 50 for a call, below 50 for a put."
    if "EMA 20" in name:
        return f"{who.capitalize()} above its 20-minute EMA for a call, below for a put."
    if "Supertrend" in name:
        return f"Supertrend (10, 3) of {who} up for a call, down for a put."
    if "09:15 open" in name:
        return f"{who.capitalize()} above its 09:15 open for a call, below for a put."
    if "VWAP" in name:
        return f"{who.capitalize()} above its VWAP for a call, below for a put."
    if "yesterday" in name:
        return "Nifty back above yesterday's close (the gap filled) for a call, still below it for a put."
    return ""


def lines(B: dict, rows: pd.DataFrame, sel: pd.DataFrame, info: dict) -> list[tuple[str, object]]:
    b = info["base"]
    out = [("S1 v12 fade · more entry combinations on 'Nifty up yesterday + gap down' days (option, stock, Nifty, VIX, clock)", FB),
           (f"{info['pattern trades']} pattern trades on {info['pattern days']} days. 09:16 buy: fail {b['fail %']:.0f} %, mean {b['mean']:+.2f} %, "
            f"median {b['median']:+.2f} %, total ₹{b['rs']:,.0f} (older half {b['older']:+.2f} %, newer half {b['newer']:+.2f} %).", F),
           (f"{info['signs']} signs in {len(set(B['fam'].values()))} families, every 1, 2 and 3 signs with at most one per family: "
            f"{info['rules']:,} rules. Entry: the first minute 09:17 – 15:00 where the rule is true, buy at the next candle's open + 1 tick, "
            "v12 exit (−20 % stop, 2 % trail once +10 %, else 15:19), real charges. 'With the trade' = up for a call, down for a put.", F),
           ("", F), ("Pass ladder", FB),
           (f"1. Enough trades (≥ 20, ≥ 8 in each half): {info['c1']:,} rules", F),
           (f"2. + the rule's own trades make money in both halves: {info['c2']:,}", F),
           (f"3. + as a whole strategy (no sign = no trade) it beats the 09:16 buy in both halves ({b['older']:+.2f} % older, "
            f"{b['newer']:+.2f} % newer, per pattern trade): {info['c3']:,}", F),
           (f"4. + the rule's own net % has a date-clustered t ≥ 2: {info['c4']:,} (clock-only rules among them: {info['clock pass']})", F),
           (f"5. + beats the shuffles (family-wise, all {info['rules']:,} rules paid for): {info['c5']:,}", F),
           (f"Luck check for steps 1 – 4: with each rule's entry times shuffled between trades of the same half (timing kept, chart reading "
            f"removed), on average {info['luck mean']:.0f} rules still pass (95 % of shuffles: up to {info['luck p95']:.0f}). The real data gives "
            f"{info['c4'] - info['clock pass']} (not counting clock-only rules); {info['luck p'] * 100:.1f} % of shuffles reach that many. "
            + "This count is noisy (heavily overlapping rules), so read it as 'about what luck gives' unless it is far out in the tail."
            + (" No single rule stands out from 15,000 tries (step 5)." if info["c5"] == 0 else ""), F),
           ("", F), ("Why the bar changed after the first run", FB),
           (f"The first run compared each rule with the 09:16 buy on the SAME trades. {info['old 4']:,} rules passed that way, "
            f"{info['old 4 with a dip sign %']:.0f} % of them needing 'option dipped 20 / 30 % below the 09:16 buy'. On those trades the 09:16 buy "
            f"had already hit its −20 % stop (it averaged {info['old 4 same trades at 09:16']:+.1f} % there), so beating it was nearly automatic. "
            "The fair question is whether the rule's own trades make money and whether the rule beats the 09:16 buy as a whole strategy. "
            "The old-bar columns stay in the Rules sheet.", F),
           ("", F), ("Top rules (furthest up the ladder first, then by the net t)", FB)]
    for _, r in rows.head(12).iterrows():
        out.append((f"#{r['rank']} {r['entry rule']}: {int(r['trades taken'])} trades, median entry {r['median entry time']}, fail {r['fail %']:.0f} %, "
                    f"mean {r['mean net %']:+.2f} % (t {r['net t (date-clustered)']:.2f}), median {r['median net %']:+.1f} %, ₹{r['total Rs']:,.0f}; "
                    f"older half {r['older half mean net %']:+.2f} % a trade ({r['older half per pattern trade %']:+.2f} % per pattern trade), newer "
                    f"{r['newer half mean net %']:+.2f} % ({r['newer half per pattern trade %']:+.2f} %); other days {r['other days: mean net %']:+.2f} %; "
                    f"median max up while held {r['median max up while held %']:.0f} %; family-wise shuffle p "
                    f"{r['shuffle p (all rules, family-wise)']:.2f}; passes all: {r['passes all']}", F))
    old = rows[rows["entry rule"] == "opt RSI back above 30 (after < 30) + opt above EMA 50"]
    if len(old):
        r = old.iloc[0]
        out += [("", F), (f"Last study's best rule (RSI back above 30 + above EMA 50) comes out the same here: {int(r['trades taken'])} trades, mean "
                          f"{r['mean net %']:+.2f} %, ₹{r['total Rs']:,.0f}; net t {r['net t (date-clustered)']:.2f}, rank #{r['rank']:,}.", F)]
    out += [("", F), ("Picked on one half, scored on the other", FB)]
    for nm in ("older half", "newer half"):
        s = sel[sel["picked on"] == nm]
        if len(s):
            out.append((f"Top {len(s)} rules by mean net % on the {nm} only: {int((s['still makes money there'] == 'yes').sum())} of {len(s)} still make "
                        f"money on the other half (average there {s['mean net % on the other half'].mean():+.2f} %). Of all eligible rules, "
                        f"{s['all eligible rules that make money on that half %'].iloc[0]:.0f} % make money on that half anyway.", F))
    mp = info["missing pattern"]
    out += [("", F), (f"Data: stock / Nifty / VIX 1-minute candles missing on {len(mp)} pattern trade(s) ({'; '.join(mp) or 'none'}); "
                      "a sign with no data never fires there. Research only: nothing live changes.", F),
            ("Sheets: Pass ladder · Rules (all; green = passes all, yellow = passes steps 1 – 4) · Signs · Picked on one half · "
             "Top 20 trades (#rank from the Rules sheet).", F)]
    return out


def write(path, B: dict, rows: pd.DataFrame, sel: pd.DataFrame, info: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    for i, (txt, font) in enumerate(lines(B, rows, sel, info), 1):
        c = ws.cell(row=i, column=1, value=txt)
        c.font = font
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 170
    b = info["base"]
    lad = pd.DataFrame({"step": [1, 2, 3, 4, 5], "what it asks": [
        "At least 20 trades taken, at least 8 in each half", "The rule's own trades make money in BOTH halves (mean net % > 0)",
        f"As a whole strategy (no sign = no trade = 0) it beats the 09:16 buy in BOTH halves: above {b['older']:+.2f} % (older) and "
        f"{b['newer']:+.2f} % (newer) per pattern trade",
        "The rule's own net % has a date-clustered t of at least 2",
        "Its t beats the best shuffled t of all rules in 95 % of 600 shuffles (entry times swapped between trades of the same half)"],
        "rules left": [info["c1"], info["c2"], info["c3"], info["c4"], info["c5"]]})
    lad["rules left"] = lad["rules left"].astype(int)
    sheet(wb, "Pass ladder", lad, f"Each step keeps only the rules that passed the step before. Shuffled data passes steps 1 – 4 with "
          f"{info['luck mean']:.0f} rules on average (95 %: {info['luck p95']:.0f}).")
    wb["Pass ladder"].column_dimensions["B"].width = 110
    out = rows.round(3)
    sheet(wb, "Rules", out, "Every rule on the pattern days (other-days columns = the same rule on every other trade, as a control). "
          "Green = passes all five steps, yellow = passes steps 1 – 4. 'Old bar' columns = the same-trade comparison the first run used. Filter any column.")
    w = wb["Rules"]
    p_all, p3 = out["passes all"].to_numpy() == "yes", out["4 net t ≥ 2"].to_numpy() == "yes"
    for i in np.flatnonzero(p_all | p3):
        for j in range(1, len(out.columns) + 1):
            w.cell(row=i + 3, column=j).fill = GOOD if p_all[i] else NEAR
    one = out[out["signs"] == 1].set_index("entry rule")
    sg = pd.DataFrame({"sign": B["names"], "family": [B["fam"][n] for n in B["names"]], "what it means": [means(n) for n in B["names"]]})
    for col in ("trades taken", "median entry time", "mean net %", "net t (date-clustered)", "per pattern trade % (no sign = no trade)", "passes all"):
        sg[f"alone: {col}"] = sg["sign"].map(one[col])
    sheet(wb, "Signs", sg, "Every sign used, what it means, and how it does on its own. 'With the trade'" + WITH + ".")
    wb["Signs"].column_dimensions["C"].width = 90
    sheet(wb, "Picked on one half", sel, "The top 10 rules by mean net % on one half only, then the same rules on the half they were not "
          "picked on. If picking worked, they would do clearly better there than the share of all rules that make money anyway.")
    f = B["f"]
    pt = np.flatnonzero(f["pattern"].to_numpy())
    tr = f.loc[pt, ["day0", "half", "symbol", "side", "09:16 net %"]].reset_index(drop=True)
    tr["half"] = tr["half"].map({"train": "older", "test": "newer"})
    ri = {r: i for i, r in enumerate(" + ".join(B["names"][q] for q in x) for x in B["rules"])}
    for _, r in rows.head(20).iterrows():
        i = ri[r["entry rule"]]
        J = B["J"][i, pt]
        cl = B["clock"][i, pt]
        tr[f"#{r['rank']} at"] = [None if c == 376 else f"{9 + (c + 15) // 60:02d}:{(c + 15) % 60:02d}" for c in cl]
        tr[f"#{r['rank']} net %"] = B["tab"]["net %"][pt, J]
    sheet(wb, "Top 20 trades", tr, "Each pattern trade: the 09:16 result and, for the top 20 rules (#rank from the Rules sheet), the entry "
          "time and net %. Blank = no sign by 15:00.")
    wb.save(path)
