"""S1 v10 fade · premium-floor sweep as a workbook to check by hand: every trade at the lowest floor (one tick) with its
day-0 option open / high / low / close, and a summary of live formulas with the floor as an input cell.

    python -m scanner.fade_floor_sheet      # after scanner.fade_floor_sweep's runs exist

Reads data/results/S1_v10_trades_floor0p05.csv and _floor2.csv and the cached fo_prices_YYYYMMDD.csv files; writes
data/results/S1_v10_floor_sweep.xlsx. A trade is taken at floor f when its fill >= f: the walk is otherwise the same at
every floor (checked: the Rs 2 run's trades are identical in the one-tick run), so one trade list covers all floors.
"""
from __future__ import annotations

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .backtest_s1_v10 import quotes_of
from .build import ROOT

RES, CACHE = ROOT / "data" / "results", ROOT / "data" / "cache"
F = "Arial"
HEAD = Font(name=F, bold=True, color="FFFFFF")
HFILL = PatternFill("solid", fgColor="1F3864")
INPUT = Font(name=F, color="0000FF", bold=True)
YELLOW = PatternFill("solid", fgColor="FFFF00")
BODY, BOLD = Font(name=F), Font(name=F, bold=True)
THIN = Border(bottom=Side(style="thin", color="BFBFBF"))

COLS = [  # (header, source column or None for a formula, number format, width)
    ("Symbol", "symbol", "@", 12), ("Day 0", "day0", "@", 11), ("Filed", "filed", "@", 11), ("Side", "side", "@", 6),
    ("NP YoY %", "np_yoy", "0.0", 9), ("NP QoQ %", "np_qoq", "0.0", 9), ("Rev YoY %", "rev_yoy", "0.0", 9),
    ("Stock prev close", "spot_prev_close", "#,##0.00", 12), ("Expiry", "expiry", "@", 11), ("Strike", "strike", "#,##0.00", 11),
    ("Strikes walked", "strikes_walked", "0", 8), ("Lot", "lot", "#,##0", 8), ("Contracts day 0", "contracts_day0", "#,##0", 10),
    ("Opt open", "o", "0.00", 8), ("Opt high", "h", "0.00", 8), ("Opt low", "l", "0.00", 8), ("Opt close", "c", "0.00", 8),
    ("Entry (open+1 tick)", "entry_price", "0.00", 10), ("E1 exit (close-1 tick)", "e1_exit_price", "0.00", 10),
    ("Lot cost Rs", None, "#,##0", 10), ("E1 net Rs (after costs)", "e1_net", "#,##0;[Red]-#,##0", 12),
    ("E1 net %", None, "0.0%;[Red]-0.0%", 9), ("E1 gross % (no costs)", None, "0.0%;[Red]-0.0%", 10),
    ("Day-high gross % (hindsight)", None, "0.0%;[Red]-0.0%", 11),
    ("E2 exit date", "e2_exit_date", "@", 11), ("E2 reason", "e2_reason", "@", 9), ("E2 net Rs", "e2_net", "#,##0;[Red]-#,##0", 10),
    ("Stock day-0 move vs headline %", "stock_oc_signed_pct", "0.00", 11), ("In v10 Rs 2 run", "in_v10", "@", 8),
    ("Fill bucket", None, "@", 10), ("Include (1/0)", "include", "0", 9),
]
C = {h: get_column_letter(i + 1) for i, (h, *_) in enumerate(COLS)}


def ohlc(t: pd.DataFrame) -> pd.DataFrame:
    """Day 0's open / high / low / close of each traded contract, from the cached option prices."""
    out = []
    for d0, g in t.groupby("day0"):
        p = pd.read_csv(CACHE / f"fo_prices_{d0.replace('-', '')}.csv")
        for i, r in g.iterrows():
            q = quotes_of(p, r["symbol"], r["expiry"], r["side"]).get(float(r["strike"]))
            out.append((i, *((q.open, q.high, q.low, q.close) if q else (None,) * 4)))
    return pd.DataFrame(out, columns=["i", "o", "h", "l", "c"]).set_index("i")


def trades() -> pd.DataFrame:
    t = pd.read_csv(RES / "S1_v10_trades_floor0p05.csv")
    b = pd.read_csv(RES / "S1_v10_trades_floor2.csv")
    t, b = t[t["entry"] == "open"].copy(), b[b["entry"] == "open"]
    k = ["symbol", "day0", "side"]
    t["in_v10"] = t.set_index(k).index.isin(b.set_index(k).index)
    t["in_v10"] = t["in_v10"].map({True: "yes", False: "no"})
    t = t[t["e1_excluded"].fillna("") == ""].join(ohlc(t))   # E1-valid trades, as the sweep's E1 table
    t["include"] = 1
    return t.sort_values(["day0", "symbol"]).reset_index(drop=True)


def write_trades(ws, t: pd.DataFrame) -> int:
    for j, (h, _, _, w) in enumerate(COLS, 1):
        c = ws.cell(1, j, h)
        c.font, c.fill, c.alignment = HEAD, HFILL, Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[1].height = 45
    for i, r in enumerate(t.itertuples(index=False), 2):
        row = r._asdict()
        f = {"Lot cost Rs": f"={C['Entry (open+1 tick)']}{i}*{C['Lot']}{i}",
             "E1 net %": f"={C['E1 net Rs (after costs)']}{i}/{C['Lot cost Rs']}{i}",
             "E1 gross % (no costs)": f"={C['E1 exit (close-1 tick)']}{i}/{C['Entry (open+1 tick)']}{i}-1",
             "Day-high gross % (hindsight)": f'=IF({C["Opt high"]}{i}="","",{C["Opt high"]}{i}/{C["Entry (open+1 tick)"]}{i}-1)',
             "Fill bucket": f'=IF({C["Entry (open+1 tick)"]}{i}<1,"under Rs 1",IF({C["Entry (open+1 tick)"]}{i}<2,"Rs 1-2","Rs 2+"))'}
        for j, (h, src, fmt, _) in enumerate(COLS, 1):
            v = f[h] if src is None else row.get(src)
            if isinstance(v, float) and v != v:
                v = None
            c = ws.cell(i, j, v)
            c.number_format, c.font = fmt, (INPUT if h == "Include (1/0)" else BODY)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{len(t) + 1}"
    ws[f"{C['Include (1/0)']}1"].comment = Comment("Set 0 to leave a trade out of every Summary formula "
                                                   "(e.g. the KALYANKJIL 2025-08-08 outlier).", "Claude")
    ws[f"{C['Day-high gross % (hindsight)']}1"].comment = Comment(
        "Selling at the day's high is not a rule: nobody knows the high in advance. Shown only to compare the "
        "best case with the 15:20 exit. Gross: before charges and slippage.", "Claude")
    return len(t) + 1


def write_summary(ws, last: int) -> None:
    rng = lambda h: f"Trades!${C[h]}$2:${C[h]}${last}"                  # noqa: E731
    fill, inc, net, cost = rng("Entry (open+1 tick)"), rng("Include (1/0)"), rng("E1 net Rs (after costs)"), rng("Lot cost Rs")
    pct, hi, e2 = rng("E1 net %"), rng("Day-high gross % (hindsight)"), rng("E2 net Rs")
    ws.column_dimensions["A"].width = 44
    for col in "BCDEFGHI":
        ws.column_dimensions[col].width = 14
    ws["A1"] = "S1 v10 fade · premium-floor backtest, 2025-01-01 to 2026-09-25"
    ws["A1"].font = Font(name=F, bold=True, size=14)
    ws["A2"] = ("Entry at day 0's option open + 1 tick, sold at day 0's close - 1 tick (E1, v10's same-day exit). Net after "
                "charges and slippage. A trade counts at floor f when its fill >= f. Every number below is a formula on "
                "the Trades sheet: change the yellow cell or the Include column and it recalculates.")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:I2")
    ws.row_dimensions[2].height = 48
    ws["A4"], ws["B4"] = "Floor to test (Rs)", 1
    ws["A4"].font, ws["B4"].font, ws["B4"].fill = BOLD, INPUT, YELLOW
    ws["B4"].comment = Comment("Input: the premium floor. v10 today uses Rs 2.", "Claude")

    hdr = ["Trades", "N", "Wins", "Hit rate", "Mean net %", "Total net Rs", "Avg lot Rs", "Mean day-high % (hindsight)", "E2 total net Rs"]

    def block(top: int, title: str, rows: list[tuple[str, str]]) -> None:
        ws.cell(top, 1, title).font = Font(name=F, bold=True, size=12)
        for j, h in enumerate(hdr, 1):
            c = ws.cell(top + 1, j, h)
            c.font, c.fill, c.alignment = HEAD, HFILL, Alignment(wrap_text=True)
        for i, (label, cond) in enumerate(rows, top + 2):
            n = f"COUNTIFS({inc},1{cond})"
            vals = [label, f"={n}", f'=COUNTIFS({inc},1,{net},">0"{cond})', f"=IF(C{i}=0,\"-\",C{i}/B{i})",
                    f"=IF(B{i}=0,\"-\",SUMIFS({pct},{inc},1{cond})/B{i})", f"=SUMIFS({net},{inc},1{cond})",
                    f"=IF(B{i}=0,\"-\",SUMIFS({cost},{inc},1{cond})/B{i})",
                    f'=IF(B{i}=0,"-",SUMIFS({hi},{inc},1{cond})/B{i})', f"=SUMIFS({e2},{inc},1{cond})"]
            for j, v in enumerate(vals, 1):
                c = ws.cell(i, j, v)
                c.font, c.border = BODY, THIN
                c.number_format = {4: "0.0%", 5: "0.0%;[Red]-0.0%", 6: "#,##0;[Red]-#,##0", 7: "#,##0",
                                   8: "0.0%;[Red]-0.0%", 9: "#,##0;[Red]-#,##0"}.get(j, "General")

    block(6, "At the floor in B4", [
        ("All trades taken (fill >= floor)", f',{fill},">="&$B$4'),
        ("  of which v10 already takes (fill >= Rs 2)", f',{fill},">=2"'),
        ("  of which the lower floor adds (floor <= fill < Rs 2)", f',{fill},">="&$B$4,{fill},"<2"'),
    ])
    block(12, "Every floor side by side (the added trades: floor <= fill < Rs 2)", [
        (f"Floor Rs {f:g}: added trades", f',{fill},">={f}",{fill},"<2"') for f in (1.5, 1.0, 0.5, 0.25, 0.05)
    ] + [("v10 (Rs 2): all trades", f',{fill},">=2"')])
    block(22, "By fill size", [
        ("Fill under Rs 1", f',{fill},"<1"'), ("Fill Rs 1 to 2", f',{fill},">=1",{fill},"<2"'),
        ("Fill Rs 2 to 5", f',{fill},">=2",{fill},"<5"'), ("Fill Rs 5 and up", f',{fill},">=5"'),
    ])
    notes = [
        "How to read it",
        "- \"Mean net %\" is the average of each trade's return; one big winner can lift it a lot (KALYANKJIL 2025-08-08 "
        "made +3,267 %). \"Total net Rs\" and \"Hit rate\" show whether most trades made money.",
        "- To see a result without one trade, set its Include cell (Trades sheet, last column) to 0.",
        "- \"Day-high %\" is hindsight: the best price of the day, which no rule can be sure to sell at. It is not the backtest.",
        "- E2 = the carried-trade rule (stop -20 %, trail 25 %, or T-1), for reference.",
        "- Limits: no order book in daily data. The fill is the open + 1 tick and the exit the close - 1 tick, and for a "
        "Rs 1 option one tick is 5 %. Not out of sample (S1 v8 used 2025-26). Pre-registered: "
        "strategies/S1_v10_backtest_2025_26.md, addendum 2026-10-19.",
    ]
    for i, s in enumerate(notes, 29):
        c = ws.cell(i, 1, s)
        c.font = BOLD if i == 29 else BODY
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=9)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[i].height = 15 if i == 29 else 30


def main() -> None:
    t = trades()
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    tr = wb.create_sheet("Trades")
    last = write_trades(tr, t)
    write_summary(ws, last)
    added = t[t["entry_price"] < 2]
    ad = wb.create_sheet("Added under Rs 2")
    ad.append(["The trades a floor below Rs 2 adds, worst to best (values; the Trades sheet has the same rows)."])
    cols = ["symbol", "day0", "side", "strike", "lot", "o", "h", "l", "c", "entry_price", "e1_exit_price", "e1_net", "e1_net_pct",
            "stock_oc_signed_pct"]
    ad.append(["Symbol", "Day 0", "Side", "Strike", "Lot", "Opt open", "Opt high", "Opt low", "Opt close", "Entry", "E1 exit",
               "E1 net Rs", "E1 net %", "Stock move vs headline %"])
    for r in added.sort_values("e1_net_pct")[cols].itertuples(index=False):
        ad.append(list(r))
    for row in ad.iter_rows():
        for c in row:
            c.font = BODY
    for c in ad[2]:
        c.font, c.fill = HEAD, HFILL
    for j in range(1, 15):
        ad.column_dimensions[get_column_letter(j)].width = 12
    ad.freeze_panes = "A3"
    wb.calculation.fullCalcOnLoad = True
    out = RES / "S1_v10_floor_sweep.xlsx"
    wb.save(out)
    print(out, len(t), "trades,", len(added), "added under Rs 2; missing OHLC:", int(t["o"].isna().sum()))


if __name__ == "__main__":
    main()
