"""S1_armA_peak_study_rows.csv -> an Excel workbook: per trade the result quality (QoQ / YoY), the option's peak
and what each alternative exit would have made; summary sheets by result category and by exit rule.

    python -m scanner.peak_study_sheet
"""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .build import ROOT
from .peak_study import rules
from .stats import clustered_t

RES = ROOT / "data" / "results"
RULE_NAME = {"live": "Live rules", "target:0.25": "Target +25%", "target:0.5": "Target +50%", "target:1.0": "Target +100%",
             "trail:0.3": "Trail 30% from peak", "trail:0.4": "Trail 40%", "trail:0.5": "Trail 50%",
             "fixed:2": "Sell day 2", "fixed:3": "Sell day 3", "fixed:5": "Sell day 5",
             "combo:0.5:0.3": "Target 50% or trail 30%", "combo:0.5:0.4": "Target 50% or trail 40%",
             "combo:1.0:0.3": "Target 100% or trail 30%", "combo:1.0:0.4": "Target 100% or trail 40%",
             "half:0.5:0.4": "Half at +50%, rest trails 40%", "half:1.0:0.4": "Half at +100%, rest trails 40%"}
HEAD = PatternFill("solid", fgColor="1F2937")
SUB = PatternFill("solid", fgColor="E2E8F0")
STRIPE = PatternFill("solid", fgColor="F8FAFC")
THIN = Side(style="thin", color="CBD5E1")
BASE, BOLD = Font(name="Calibri", size=10), Font(name="Calibri", size=10, bold=True)
PCT = '+0.0"%";-0.0"%";0.0"%"'


def signed(v, bold=False) -> Font:
    return Font(name="Calibri", size=10, bold=bold, color="15803D" if v > 0 else "B91C1C" if v < 0 else "000000")


def write_table(ws, df: pd.DataFrame, pct_cols: set[str], start_row: int = 1, title: str | None = None) -> int:
    r = start_row
    if title:
        ws.cell(row=r, column=1, value=title).font = Font(bold=True, size=12)
        r += 1
    for j, c in enumerate(df.columns, 1):
        h = ws.cell(row=r, column=j, value=c)
        h.font, h.fill = Font(bold=True, color="FFFFFF", size=9), HEAD
        h.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(j)].width = max(ws.column_dimensions[get_column_letter(j)].width or 0,
                                                               min(28, max(9, len(str(c)) * 0.9)))
    ws.row_dimensions[r].height = 30
    r += 1
    for i, rec in enumerate(df.itertuples(index=False)):
        for j, (c, v) in enumerate(zip(df.columns, rec), 1):
            if v is None or (isinstance(v, float) and v != v):
                v = None
            cell = ws.cell(row=r, column=j, value=v)
            cell.font = signed(v) if c in pct_cols and isinstance(v, (int, float)) else BASE
            if c in pct_cols and isinstance(v, (int, float)):
                cell.number_format = PCT
            if i % 2:
                cell.fill = STRIPE
            cell.border = Border(bottom=THIN)
        r += 1
    return r + 1


def summary_by(df: pd.DataFrame, col: str, order=None) -> pd.DataFrame:
    out = []
    for sample in ("2025-26", "2021-24"):
        g0 = df[df["sample"] == sample]
        for val in (order or sorted(g0[col].dropna().unique(), key=str)):
            g = g0[g0[col] == val]
            if not len(g):
                continue
            out.append({"Sample": sample, col: val, "Trades": len(g),
                        "Win rate": round((g["live|net_pct"] > 0).mean() * 100, 1),
                        "Average net (live rules)": round(g["live|net_pct"].mean(), 1),
                        "t": clustered_t(g["live|net_pct"], g["signal_date"]) if len(g) > 1 else None,
                        "Average max peak": round(g["peak_pct"].mean(), 1),
                        "Peaked >= +50%": round((g["peak_pct"] >= 50).mean() * 100, 1),
                        "Median peak day": g["peak_day"].median(),
                        "Stock to T-1 (avg)": round(g["stock_entry_to_t1_pct"].mean(), 2),
                        "Best alt. exit": max(rules(), key=lambda r: g[f"{r}|net_pct"].mean()),
                        "Its average net": round(max(g[f"{r}|net_pct"].mean() for r in rules()), 1)})
    return pd.DataFrame(out)


def exits_table(df: pd.DataFrame) -> pd.DataFrame:
    out = []
    for sample in ("2025-26", "2021-24"):
        g = df[df["sample"] == sample]
        for r in rules():
            v = g[f"{r}|net_pct"]
            out.append({"Sample": sample, "Exit rule": RULE_NAME.get(r, r), "Trades": len(g),
                        "Win rate": round((v > 0).mean() * 100, 1), "Average net": round(v.mean(), 1),
                        "Median net": round(v.median(), 1), "t": clustered_t(v, g["signal_date"]),
                        "Avg day sold": round(g[f"{r}|day"].mean(), 1),
                        "Total ₹ (1 lot each)": round((v / 100 * g["cost"]).sum())})
    return pd.DataFrame(out)


def build(out: Path = RES / "S1_armA_peak_study.xlsx") -> Path:
    df = pd.read_csv(RES / "S1_armA_peak_study_rows.csv")
    wb = Workbook()
    ws = wb.active
    ws.title = "Exit rules"
    r = write_table(ws, exits_table(df), {"Win rate", "Average net", "Median net"}, 1,
                    "What each exit rule would have made on the same trades (after charges and slippage)")
    ws.cell(row=r, column=1, value="Hindsight: every rule is applied to the same option paths; the live rules' "
                                   "numbers equal the backtest. A rule that only wins in one sample is a fit, not a rule.").font = BASE
    ws2 = wb.create_sheet("By result quality")
    r = 1
    for col, title, order in (("net_profit_qoq_cat", "Net profit vs previous quarter (QoQ)", ("up", "down", "n/a")),
                              ("net_profit_yoy_cat", "Net profit vs a year ago (YoY)", ("up", "down", "n/a")),
                              ("revenue_qoq_cat", "Revenue QoQ", ("up", "down", "n/a")),
                              ("revenue_yoy_cat", "Revenue YoY", ("up", "down", "n/a")),
                              ("eps_yoy_cat", "EPS YoY", ("up", "down", "n/a")),
                              ("beat_cat", "Results-day beat vs sector", ("<3%", "3-5%", ">5%")),
                              ("turned_to_loss", "Profit turned into a loss this quarter", (False, True))):
        r = write_table(ws2, summary_by(df, col, order), {"Win rate", "Average net (live rules)", "Average max peak",
                                                           "Peaked >= +50%", "Its average net"}, r, title)
    ws3 = wb.create_sheet("Trades")
    cols = ["sample", "symbol", "sector", "signal_date", "quarter", "beat_pct", "raw_move_pct",
            "revenue_qoq", "revenue_yoy", "net_profit_qoq", "net_profit_yoy", "eps_qoq", "eps_yoy", "turned_to_loss",
            "entry_date", "strike", "distance_pct", "entry_price", "cost", "peak_pct", "peak_day", "dd_after_peak_pct",
            "first_day_50", "first_day_100", "end_pct", "stock_entry_to_peak_pct", "stock_entry_to_t1_pct", "stock_max_pct",
            "live_exit_date", "live_reason"] + [f"{r}|net_pct" for r in rules()]
    t = df[cols].rename(columns={**{f"{r}|net_pct": RULE_NAME.get(r, r) for r in rules()},
                                 "beat_pct": "beat vs sector %", "raw_move_pct": "day-0 move %",
                                 "peak_pct": "max peak %", "dd_after_peak_pct": "fall after peak %",
                                 "end_pct": "at T-1 %", "stock_entry_to_peak_pct": "stock at peak day %",
                                 "stock_entry_to_t1_pct": "stock at T-1 %", "stock_max_pct": "stock max %"})
    pct_cols = {c for c in t.columns if "%" in c or c in RULE_NAME.values() or c.endswith("_qoq") or c.endswith("_yoy")}
    write_table(ws3, t, pct_cols, 1)
    ws3.freeze_panes = "C2"
    ws3.auto_filter.ref = ws3.dimensions
    wb.save(out)
    return out


if __name__ == "__main__":
    print(build())
