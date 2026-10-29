"""Excel writer for scanner.fade_market_pattern: Summary first, pattern tables, then models A (09:16) and B (whole day)."""
from __future__ import annotations

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from .fade_indicators_xlsx import F, FB, sheet

COLS = ["test", "method", "trades", "AUC", "flagged trades", "fail % flagged", "fail % kept", "fail % all",
        "mean net % flagged", "mean net % kept", "mean net % all", "total Rs kept (skip flagged)", "total Rs all"]


def model_lines(R: dict) -> list[str]:
    out = []
    for df in (R["halves"], R["wf"]):
        r = df.set_index("method")
        rule = r.iloc[0]
        out.append(f"{df['test'].iloc[0]}: {r.index[0].replace('Rule: ', 'rule ')} AUC {rule['AUC']:.2f}, keeps {rule['mean net % kept']:+.2f} % a trade "
                   f"(all {rule['mean net % all']:+.2f} %)")
        for m in ("Logistic regression", "Gradient boosting", "Small decision tree"):
            a = r.loc[m]
            better = a["AUC"] > rule["AUC"] and a["mean net % kept"] > rule["mean net % kept"]
            out.append(f"    {m}: AUC {a['AUC']:.2f}, skipping its {int(a['flagged trades'])} riskiest keeps {a['mean net % kept']:+.2f} % "
                       f"→ {'beats the rule' if better else 'does not beat the rule'}; "
                       f"{'better' if a['mean net % kept'] > a['mean net % all'] else 'worse'} than taking every trade")
    return out


def pattern_lines(pat: pd.DataFrame) -> list[str]:
    s = pat[(pat["same way both halves"] == "yes") & (pat["p (fail vs rest)"] < 0.05)].sort_values("p (fail vs rest)")
    if s.empty:
        return ["No market path passes (p < 0.05 and the same way in both halves)."]
    return [f"{r['group']} → {r['condition']}: {r['trades']} trades, fail {r['fail %']:.0f} % vs {r['fail % rest']:.0f} % for the rest "
            f"(older half {r['train fail %']:.0f} %, newer half {r['test fail %']:.0f} %; p {r['p (fail vs rest)']:.3f}); {r['reads as']}"
            for _, r in s.iterrows()]


def tree_lines(R: dict, tag: str) -> list[str]:
    lv = R["leaves"]
    held = lv[(lv["held in the test half"] == "yes") & (lv["train fail %"] > 50) & (lv["test trades"] >= 10)]
    if held.empty:
        return [f"{tag}: no fail-heavy box of the tree held on the newer half."]
    return [f"{tag}: {r['pattern']} → fail {r['train fail %']:.0f} % older half ({r['train trades']}), {r['test fail %']:.0f} % newer half ({r['test trades']})"
            for _, r in held.iterrows()]


def write(path, f, pat, A, B) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    lines = [("S1 v12 fade · market only (no RSI / EMA / Supertrend): what is the market doing when the fade fails?", FB),
             (f"{len(f)} backtest trades, v12 exit, real charges. Fail = net < 0 ({int(f['fail'].sum())} trades). Older half = before 3 Nov 2025, "
              "newer half = after. A pattern counts only when p < 0.05 AND it points the same way in both halves.", F),
             ("", F), ("Part 1 · market paths that pass", FB)]
    lines += [(x, F) for x in pattern_lines(pat)]
    lines += [("", F), ("Part 2 · the small tree's fail-heavy boxes that held on the newer half (learned on the older half only)", FB)]
    lines += [(x, F) for x in tree_lines(A, "A · known at 09:16") + tree_lines(B, "B · whole day")]
    lines += [("", F), ("Part 3 · can a model beat the simple rule? (scored only on trades it never saw; AUC 0.50 = coin flip)", FB)]
    lines += [(x, F) for x in model_lines(A) + model_lines(B)]
    lines += [("", F), ("A = known at 09:16 (previous day, the gaps): usable live. B = adds the move after the open: explains WHY, "
                        "but is only known after the close. Research only: nothing live changes.", F),
              ("Sheets: Market patterns · A models · A tree · A weights · B models · B tree · B weights · Trades.", F)]
    for i, (txt, font) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=txt)
        c.font = font
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 170
    sheet(wb, "Market patterns", pat, "Fail % per market path vs every other trade. Marked rows: p < 0.05 and the same way in both "
          "halves (red = fails more, green = fails less). Day path and day range are only known after the close.", "same way both halves")
    for tag, R in (("A", A), ("B", B)):
        sheet(wb, f"{tag} models", pd.concat([R["halves"], R["wf"]])[COLS], f"{tag}: rules vs the three models, on the newer half and "
              "walk-forward. Models flag as many riskiest trades as the first rule flags.")
        sheet(wb, f"{tag} tree", R["leaves"], f"{tag}: every box of the small tree learned on the older half, as a plain pattern, and "
              "what happened in that box on the newer half it never saw.")
        sheet(wb, f"{tag} weights", R["weights"], f"{tag}: logistic regression weights (older half, scaled inputs). 0 = switched off.")
    cols = ["day0", "half", "symbol", "side", "v12 net %", "v12 net Rs", "fail", "nifty prev day %", "nifty gap %", "nifty open to close %",
            "nifty close move %", "nifty high from open %", "nifty low from open %", "stock gap %", "stock prev day %"]
    tr = f[cols].copy()
    tr["fail"] = tr["fail"].map({True: "fail", False: ""})
    sheet(wb, "Trades", tr, "Every trade with its market moves. Filter any column to test your own ideas.")
    for name in wb.sheetnames[1:]:
        if "tree" in name:
            wb[name].column_dimensions["A"].width = 90
    wb.save(path)
