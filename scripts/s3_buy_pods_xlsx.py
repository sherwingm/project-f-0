"""S3 buy pods workbook for the phone: data/results/S3_buy_pods_2026-10-30.xlsx.

    python scripts/s3_buy_pods_xlsx.py

Summary first (formulas over the trade sheets), then one sheet per pod's primary trades (NIFTY F1 re-run, SENSEX R2),
the F3a gap signals and the F4 max up. Reads only data/results; run scripts/recalc afterwards for cached values.
"""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RES = Path(__file__).resolve().parents[1] / "data" / "results"
OUT = RES / "S3_buy_pods_2026-10-30.xlsx"
PRIMARY = {"E1": "E-1 → E 14:45", "Q1": "lookback 5 · pct 20 · hold 3",
           "B1": "OR 15 min · volume 2x · stop 25 %", "G1": "gap 0.5 % · confirm 15 min · stop 25 %"}
IDEA = {"E1": "Event straddle", "Q1": "Squeeze straddle", "B1": "Breakout + volume", "G1": "Gap follow-through"}
FONT, BOLD = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEAD = PatternFill("solid", fgColor="1F4E78")
PCT, RS, NUM = "0.0%;-0.0%;-", "₹#,##0;-₹#,##0;-", "0.00"


def _style_sheet(ws, widths: dict[int, int]) -> None:
    for c in ws[1]:
        c.font, c.fill = BOLD, HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = FONT
    for i, w in widths.items():
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 30


def trades_sheet(wb, title: str, t: pd.DataFrame, extra: list[str]) -> tuple[str, int]:
    ws = wb.create_sheet(title)
    cols = ["Day", "Exit day", "Exit time", "Exit reason", *extra, "Legs (buy strike@in>out)", "Lot",
            "Premium paid (₹)", "Net P&L (₹)", "Net % of premium", "Worst while held"]
    ws.append(cols)
    for r in t.sort_values("day").itertuples():
        vals = [str(r.day)[:10], str(r.exit_day)[:10], r.exit_at, r.reason]
        vals += [getattr(r, e.replace(" ", "_").lower()) if hasattr(r, e.replace(" ", "_").lower()) else None for e in extra]
        vals += [r.legs, r.lot, r.premium_rs, r.net_rs, r.net_pct_premium / 100, r.worst_pct / 100]
        ws.append(vals)
    n = ws.max_row
    fmt = {len(cols) - 3: RS, len(cols) - 2: RS, len(cols) - 1: PCT, len(cols): PCT}
    for row in ws.iter_rows(min_row=2):
        for i, f in fmt.items():
            row[i - 1].number_format = f
    _style_sheet(ws, {1: 11, 2: 11, 3: 9, 4: 10, **{5 + i: 12 for i in range(len(extra))},
                      5 + len(extra): 42, 6 + len(extra): 6, 7 + len(extra): 12, 8 + len(extra): 12,
                      9 + len(extra): 11, 10 + len(extra): 11})
    return title, n


def primary(name: str) -> pd.DataFrame | None:
    p = RES / f"S3_{name}_trades.csv"
    if not p.exists():
        return None
    t = pd.read_csv(p)
    pod = name.split("_")[-1]
    return t[(t["cell"] == PRIMARY[pod]) & (t["cost_mult"] == 1.0) & (t["skip"].fillna("") == "")]


def main() -> None:
    wb = Workbook()
    summ = wb.active
    summ.title = "Summary"
    summ.append(["Market", "Pod", "Idea", "Trades", "Won", "Win rate", "Mean net % of premium", "Median",
                 "t (mean ÷ s.e.)", "Total net (₹)", "Gate verdict (from the report)", "Trades sheet"])
    sheets = []
    for market, prefix in (("NIFTY", ""), ("SENSEX", "R2_")):
        for pod in PRIMARY:
            t = primary(prefix + pod)
            if t is None:
                continue
            extra = {"E1": ["Event type"], "B1": ["Side", "Signal"], "G1": ["Side", "Gap pct"]}.get(pod, [])
            if pod == "E1":
                t = t.rename(columns={"event_type": "event_type"})
            title, n = trades_sheet(wb, f"{market} {pod}", t, extra)
            rep = (RES / f"S3_{prefix}{pod}_report.md").read_text(encoding="utf-8")
            verdict = "PASS" if "**PASS" in rep or "Verdict: PASS" in rep else "FAIL"
            sheets.append((market, pod, title, n, len(extra), verdict))
    for i, (market, pod, title, n, ne, verdict) in enumerate(sheets, start=2):
        q = f"'{title}'"
        pc = get_column_letter(9 + ne)          # Net % of premium
        rc = get_column_letter(8 + ne)          # Net P&L (₹)
        rng, rrs = f"{q}!{pc}2:{pc}{n}", f"{q}!{rc}2:{rc}{n}"
        summ.append([market, pod, IDEA[pod], f"=COUNT({rng})", f'=COUNTIF({rrs},">0")', f"=IF(D{i}>0,E{i}/D{i},0)",
                     f"=IF(D{i}>0,AVERAGE({rng}),0)", f"=IF(D{i}>0,MEDIAN({rng}),0)",
                     f"=IF(D{i}>1,AVERAGE({rng})/(STDEV({rng})/SQRT(D{i})),0)", f"=SUM({rrs})", verdict, title])
    for row in summ.iter_rows(min_row=2):
        row[5].number_format = row[6].number_format = row[7].number_format = PCT
        row[8].number_format, row[9].number_format = NUM, RS
    note = summ.max_row + 2
    summ.cell(note, 1, "Notes").font = Font(name="Arial", size=10, bold=True)
    for k, s in enumerate([
        "Window 2024-10-01 → 2026-10-28. NIFTY = the data-fix re-run (spot store 62a23257); SENSEX = R2, lot fixed at 20.",
        "% = net result after costs, as a share of the premium paid. Gate: ≥ 30 trades, mean > 0 also at 2× costs, t ≥ 2,",
        "both halves positive, worst trade smaller than the average month, neighbours positive (S3 rules, DECISIONS.md).",
        "Rules pre-registered: strategies/S3_E1/Q1/B1/G1 pages and strategies/S3_buy_followups_2026-10-30.md.",
    ], start=1):
        summ.cell(note + k, 1, s).font = FONT
    _style_sheet(summ, {1: 9, 2: 6, 3: 20, 4: 8, 5: 7, 6: 9, 7: 12, 8: 10, 9: 10, 10: 13, 11: 14, 12: 12})
    summ.auto_filter.ref = f"A1:L{len(sheets) + 1}"

    f3 = RES / "S3_F3a_gap_spot_signals.csv"
    if f3.exists():
        s = pd.read_csv(f3)
        ws = wb.create_sheet("Gap check 2022-24")
        ws.append(["Day", "Gap side", "Gap %", "Entry 09:30", "Close 15:15", "Follow-through to 15:15",
                   "Max favourable", "Max adverse"])
        for r in s.itertuples():
            ws.append([r.day, r.side, r.gap_pct / 100, r.entry, r.exit_1515, r.follow_pct / 100, r.max_fav_pct / 100,
                       r.max_adv_pct / 100])
        for row in ws.iter_rows(min_row=2):
            for i in (2, 5, 6, 7):
                row[i].number_format = "0.00%"
            row[3].number_format = row[4].number_format = "#,##0.00"
        _style_sheet(ws, {1: 11, 2: 9, 3: 8, 4: 11, 5: 11, 6: 13, 7: 12, 8: 12})
    f4 = RES / "S3_F4_maxup_trades.csv"
    if f4.exists():
        m = pd.read_csv(f4)
        ws = wb.create_sheet("Max up NIFTY")
        ws.append(["Pod", "Day", "Entry", "Exit day", "Exit time", "Exit reason", "Booked net % of premium",
                   "Max up while held", "Peak while held at", "Max up any time (to 15:29)", "Peak any time at"])
        for r in m.itertuples():
            ws.append([r.pod, r.day, r.entry, r.exit_day, r.exit_at, r.reason, r.net_pct_premium / 100,
                       r.max_up_held_pct / 100, r.peak_held_at, r.max_up_any_pct / 100, r.peak_any_at])
        for row in ws.iter_rows(min_row=2):
            for i in (6, 7, 9):
                row[i].number_format = PCT
        _style_sheet(ws, {1: 6, 2: 11, 3: 7, 4: 11, 5: 8, 6: 9, 7: 12, 8: 12, 9: 17, 10: 14, 11: 17})
    wb.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
